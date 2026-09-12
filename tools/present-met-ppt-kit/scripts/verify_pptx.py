#!/usr/bin/env python3
"""PPTX 结构校验:页数、比例、空页和可选图片计数。"""
import argparse
import sys
from pptx import Presentation

EMU_PER_INCH = 914400
DEFAULT_TOLERANCE = 0.02


def parse_aspect(value):
    try:
        width, height = (float(part) for part in value.split(":", 1))
        if width <= 0 or height <= 0:
            raise ValueError
        return width / height
    except (ValueError, TypeError):
        raise argparse.ArgumentTypeError("aspect 必须形如 16:9")


def parse_picture_counts(value):
    try:
        counts = [int(part) for part in value.split(",")]
        if not counts or any(count < 0 for count in counts):
            raise ValueError
        return counts
    except (ValueError, TypeError):
        raise argparse.ArgumentTypeError("picture-counts 必须是逗号分隔的非负整数")


def build_parser():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pptx")
    parser.add_argument("--expected-count", type=int)
    parser.add_argument("--aspect", type=parse_aspect)
    parser.add_argument("--picture-counts", type=parse_picture_counts)
    parser.add_argument("--tolerance", type=float, default=DEFAULT_TOLERANCE)
    return parser


def main(argv=None):
    args = build_parser().parse_args(argv)
    if args.expected_count is not None and args.expected_count < 1:
        print("❌ expected-count 必须是正整数", file=sys.stderr)
        return 2
    if args.tolerance < 0:
        print("❌ tolerance 不能为负数", file=sys.stderr)
        return 2

    try:
        prs = Presentation(args.pptx)
    except Exception as exc:
        print(f"❌ 无法读取 PPTX: {args.pptx}: {exc}", file=sys.stderr)
        return 1

    slides = list(prs.slides)
    count = len(slides)
    width = prs.slide_width / EMU_PER_INCH
    height = prs.slide_height / EMU_PER_INCH
    print("slides:", count)
    print("size: %.3f x %.3f in" % (width, height))

    errors = []
    if count == 0:
        errors.append("PPTX 没有页面")
    if args.expected_count is not None and count != args.expected_count:
        errors.append(f"页数不符: expected {args.expected_count}, got {count}")
    if args.aspect is not None:
        actual_aspect = width / height if height else 0
        if abs(actual_aspect - args.aspect) > args.tolerance:
            errors.append(f"页面比例不符: expected {args.aspect:.6f}, got {actual_aspect:.6f}")

    empty = [i + 1 for i, slide in enumerate(slides) if len(slide.shapes) == 0]
    print("empty slides:", empty if empty else "none")
    if empty:
        errors.append("空页: " + ", ".join(str(i) for i in empty))

    picture_counts = [
        sum(1 for shape in slide.shapes if shape.shape_type == 13)
        for slide in slides
    ]
    print("picture counts:", picture_counts if picture_counts else "none")
    if args.picture_counts is not None:
        if len(args.picture_counts) != count:
            errors.append(
                f"图片计数长度不符: expected {len(args.picture_counts)}, got {count} 页"
            )
        else:
            mismatches = [
                f"P{index + 1}: expected {expected}, got {actual}"
                for index, (expected, actual) in enumerate(zip(args.picture_counts, picture_counts))
                if expected != actual
            ]
            if mismatches:
                errors.append("图片计数不符: " + "; ".join(mismatches))

    if errors:
        print("❌ 校验失败:", file=sys.stderr)
        for error in errors:
            print("   - " + error, file=sys.stderr)
        return 1
    print("✅ PPTX 结构校验通过")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
